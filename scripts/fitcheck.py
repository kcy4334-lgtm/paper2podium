# -*- coding: utf-8 -*-
"""Does the rendered deck fit in its place - overflow, overlap, text too small.

    python scripts/fitcheck.py out/talk.pdf
    python scripts/fitcheck.py out/talk.pdf --min-height 2.4 --show 8

Why this check exists on its own.
  Look across systems that build talk decks and every usable one is the same chain:
  read -> plan -> intermediate representation -> generate code -> render ->
  look and fix. The last two steps exist for the same reason every time: the most
  common failure, overflow and overlap, is invisible in the source and only
  visible in the rendered screen.

  Our chain had that spot empty. `deckcheck` reads the `.tex`, and `outcheck` reads
  the text back out of the output. But if the text arrived while covering the
  cell next to it, or printed at 6pt, both of them still pass it.

  This check reads the PDF as coordinates. It cannot replace human eyes, but it
  points at which page needs eyes on it.

What it looks at
  · Text/figures that ran off the page (this is what an overflowing table looks like)
  · Overlap between text, and between text and a figure
  · Text that is too small - measured as a ratio to page height. Measuring in
    points would mean something different for every format and could not travel.

Limit: it cannot see an ugly layout. A page that neither overlaps nor overflows
  passes even if it is ugly. That still has to be checked by eye - this check only
  tells you where to point your eyes.
"""
import argparse
import io
import os
import re
import sys

try:
    try:
        import pymupdf as fitz
    except ImportError:      # PyMuPDF before 1.24.3 has only the old import name
        import fitz
except ImportError:  # pragma: no cover
    raise SystemExit("PyMuPDF is required:  pip install pymupdf")

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    # The PDF check still runs fully without it. The only thing it cannot do is
    #   text inside figures, so it must not die here - it flags that visibly instead,
    #   when the time comes.
    Image = None

# A footnote/superscript marker is supposed to be small. Without this exclusion,
#   the two markers † and ⋆ in a real deck get flagged as "small text". A person
#   would let two slide, but the symbol repeats on every table and quickly floods
#   the list.
MARKER = re.compile(r"[\W\d]{1,2}", re.U)

# How small is too small to read - this is one of the few values you do not have
#   to invent.
#
#   AVIXA's (ANSI/AVIXA V202.01) Basic Decision Making standard:
#       cap height >= distance to the farthest seat / 200
#   Dividing this by the screen height falls straight out:
#       required cap height (%) = viewing ratio / 2
#   where **viewing ratio** = distance to the farthest seat / screen height. For a
#   classroom it is usually around 6, and the required value then is 3.0% of
#   screen height.
#
# Watch the units. That 3.0% is cap height, not font size (em). What a PDF
#   reports is font size, so it is divided by the cap ratio before comparing. The
#   ratio runs 0.62-0.72 by font, and is measured from the real glyph when possible,
#   falling back to the value below otherwise.
CAP_RATIO = 0.70
AVIXA_ACUITY = 200.0
# Whether something ran off the page is geometry. The only tolerance is one point
# of render rounding.
EPS = 0.5


def min_cap_pct(viewing_ratio):
    """Viewing ratio -> required cap height (% of screen height)."""
    return 100.0 * viewing_ratio / AVIXA_ACUITY


def cap_ratio_of(page, fallback=CAP_RATIO):
    """**Cap height / font size.** Currently not measured, a constant is used instead
    - `(value, False)`.

    Two attempts at measuring it both ended up reading something else.
      (1) `rawdict`'s glyph box is the character cell, not the glyph's ink extent,
          so uppercase and lowercase both came out 1.00.
      (2) `Font.glyph_bbox` returns the font's overall bounding box, fixed at 1.49
          regardless of the character.
      If every value comes out identical, nothing is actually being measured. Leaving
      that in place would have put 3.00% where the floor should have been 4.29%, and
      the deck would have passed.

    So this falls back to a constant, but says in the output that it fell back.
      Common fonts fall in the 0.62-0.72 range, so the error is around 10%, and a deck
      on the borderline can supply its own font's value with `--cap-ratio`. Quietly
      using an estimate lets it harden into a fact.
    """
    return fallback, False


def _hex(rgb):
    return "#%02X%02X%02X" % tuple(int(x) for x in rgb)


def spans(page):
    """(rect, text, size, text color) - empty cells are excluded."""
    out = []
    for b in page.get_text("dict")["blocks"]:
        for ln in b.get("lines", []):
            for sp in ln["spans"]:
                if sp["text"].strip():
                    out.append((fitz.Rect(sp["bbox"]), sp["text"].strip(),
                                sp["size"], sp.get("color", 0)))
    return out


MATH_MARKS = set(u"ˆ˜¯˙´`^~√′″‴¨ˇ˘→") | set(chr(c) for c in range(0x300, 0x370))
MATH_FONT = ("CMEX", "CMSY", "CMMI", "MSAM", "MSBM", "LMMath", "Math", "Symbol", "STIX",
             "rsfs", "eufm", "Cambria Math")


def math_flags(page):
    """(same order as spans()) is this piece a math accessory (accent/radical/script)
    - True/False.

    A math hat (ˆ) and the letter under it, a radical (√) and its parenthesis, and a
      subscript are naturally overlapping and naturally small. Counting them as
      overlap/small-text would mean a deck containing math could never pass fitcheck.
      Only a fragment that is short and is either a math
      symbol, in a math font, or sitting right next to a bigger character on the same
      line (a script) is excluded. Body text is checked as-is.
    """
    out = []
    for b in page.get_text("dict")["blocks"]:
        # Neighbors are looked for across the whole block - in `g_i^2` where a
        #   super/subscript sits stacked, the subscript i gets pulled out onto its own
        #   line, and a check that only looked at the same line would flag it as
        #   "small text". Only vertically overlapping spans count as neighbors.
        allsp = [s for ln in b.get("lines", []) for s in ln["spans"] if s["text"].strip()]
        for ln in b.get("lines", []):
            sps = [s for s in ln["spans"] if s["text"].strip()]
            for s in sps:
                txt = s["text"].strip()
                short = len(txt) <= 3
                mark = all(c in MATH_MARKS for c in txt)
                mfont = any(k in s.get("font", "") for k in MATH_FONT)
                r = fitz.Rect(s["bbox"])
                script = short and any(
                    o is not s and o["size"] >= s["size"] * 1.15
                    and abs(fitz.Rect(o["bbox"]).x1 - r.x0) < 3.0
                    and r.y0 < fitz.Rect(o["bbox"]).y1 and r.y1 > fitz.Rect(o["bbox"]).y0
                    for o in allsp)
                out.append(bool(mark or (short and mfont) or script))
    return out


DECK_FONTS = set()
DECK_BODY = [0.0]


def _font_key(name):
    return re.sub(r"^[A-Z]{6}\+", "", str(name or ""))      # strip a subset-font prefix (ABCDEF+)


def font_family(name):
    """The font **family** - the name with size/weight/slant stripped (`LMSans8-Oblique` ->
    `LMSans`).

    A deck's italic/bold font shows up on only a few pages, so it was not counted as
      a "deck font", and table text set in it got reported as "text inside a pasted
      paper figure" even when that deck had no pasted figures at all."""
    n = _font_key(name)
    return re.sub(r"(?:[-,]?(?:Bold|Oblique|Italic|Regular|Medium|Semibold|Light|Demi|BoldItalic|"
                  r"BoldOblique|It|Bd|BI|Slanted))+$|\d+.*$", "", n, flags=re.I) or n


def span_fonts(page):
    """Font names, in the same order as spans()."""
    out = []
    for b in page.get_text("dict")["blocks"]:
        for ln in b.get("lines", []):
            for s in ln["spans"]:
                if s["text"].strip():
                    out.append(_font_key(s.get("font")))
    return out


def deck_fonts(doc, skip=()):
    """The deck's own fonts - ones appearing on more than 30% of pages (at least two).
    A pasted figure's font only appears on that one page.

    `skip` = pages (1-based) with a pasted paper figure. Pasting a paper figure onto
      six pages pushed that figure's font over 30%, so it got counted as a "deck
      font", and the figure's own tick-mark labels then got flagged as deck text.
      Those pages are excluded from the count.
    """
    from collections import Counter
    c = Counter()
    n = 0
    for i, p in enumerate(doc, 1):
        if i in skip:
            continue
        c.update(set(span_fonts(p)))
        n += 1
    if n < 4:                     # too few pages to tell which font belongs to the deck
        return set()
    return {f for f, k in c.items() if k >= max(2, 0.3 * n)}


def images(page):
    out = []
    for info in page.get_images(full=True):
        try:
            out.append(page.get_image_bbox(info))
        except (ValueError, RuntimeError):
            pass
    return out


def _lum(rgb):
    """WCAG relative luminance. sRGB 0-255 -> 0-1."""
    out = []
    for c in rgb:
        c = c / 255.0
        out.append(c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4)
    return 0.2126 * out[0] + 0.7152 * out[1] + 0.0722 * out[2]


def contrast(a, b):
    """WCAG contrast ratio (1-21)."""
    la, lb = _lum(a), _lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


FRAME_NO = re.compile(r"\d+\s*/\s*\d+")


def paper_under(pix, rect, page_rect, ink):
    """The background color behind the text. The text color is not guessed, since
    the PDF reports it directly.

    The first version guessed both the text color and the background from
      pixels. That gave a dark-background slide's page number a contrast of `1.0:1`:
      every pixel in the box was the same color, so the "darkest 10%" and the
      "lightest 10%" came out equal, which is not physically possible. The text
      color is stored exactly in `span["color"]`. There was no reason to guess it.
    The background is taken as the most common pixel in the box that differs
      enough from the text color. A glyph's stroke only covers part of the box, so
      the mode is the background.
    """
    sx = pix.width / page_rect.width
    sy = pix.height / page_rect.height
    x0, y0 = max(0, int(rect.x0 * sx)), max(0, int(rect.y0 * sy))
    x1 = min(pix.width, int(rect.x1 * sx) + 1)
    y1 = min(pix.height, int(rect.y1 * sy) + 1)
    if x1 - x0 < 2 or y1 - y0 < 2:
        return None
    seen = {}
    for y in range(y0, y1):
        for x in range(x0, x1):
            c = pix.pixel(x, y)[:3]
            # exclude the stroke and its anti-aliased edge
            if sum(abs(c[i] - ink[i]) for i in range(3)) < 90:
                continue
            seen[c] = seen.get(c, 0) + 1
    if not seen:
        # the whole box is the text color = background and text are the same color =
        # invisible
        return ink
    return max(seen.items(), key=lambda kv: kv[1])[0]


def overlap_frac(a, b):
    """The overlapping area divided by the smaller of the two areas. 0 means no overlap."""
    r = fitz.Rect(a) & fitz.Rect(b)
    if r.is_empty:
        return 0.0
    small = min(abs(a.get_area()), abs(b.get_area()))
    return 0.0 if small <= 0 else r.get_area() / small


def fingerprint(im):
    """A fingerprint for one image - an 8x8 grayscale reduced to 16 levels. 64 hex digits.

    Size alone cannot match it. A diagram and a grid drawn at the same spot have the
      same pixel size, so a small text hit on one could get reported on the wrong
      page number, which is worse than not reporting it at all.
    """
    g = im.convert("L").resize((8, 8))
    return "".join("%x" % (p >> 4) for p in g.tobytes())


def fig_text(path):
    """`textsize.tsv` left by `build_figs` -> {fingerprint: [(pt, text, file)]}.

    Why hand this off through a file: text inside a figure does not exist in the
      PDF's text layer. Since this checker only looks at the PDF, it could not see a
      single piece of a 4.4pt diagram, and an unreadable deck passed all seven
      checkers. Only the side that drew it knows that size.
    The matching key is the **fingerprint**. The name `page.get_images()` gives is a
      PDF resource name (`Im10`), not a file name, and pixel size collides between
      figures drawn at the same spot - once that happened, page 6's caption showed up
      on pages 8 and 12 too.
    """
    out = {}
    if not path or not os.path.exists(path):
        return out
    with io.open(path, encoding="utf-8") as f_:
        for line in f_:
            if line.startswith("#") or not line.strip():
                continue
            parts = line.rstrip("\n").split("\t")
            if len(parts) < 6:
                continue
            try:
                pt = float(parts[4])
            except ValueError:
                continue
            out.setdefault(parts[3], []).append((pt, parts[5], parts[0]))
    return out


def image_prints(doc, page):
    """The fingerprints of the images placed on this page. The key for matching against
    `textsize.tsv`."""
    out = []
    for info in page.get_images(full=True):
        try:
            pix = fitz.Pixmap(doc, info[0])
            if pix.alpha or (pix.colorspace and pix.n > 3):
                pix = fitz.Pixmap(fitz.csRGB, pix)
            mode = "RGB" if pix.n >= 3 else "L"
            im = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
            out.append(fingerprint(im))
        except Exception:                                  # noqa: BLE001
            continue
    return out


SOFT_SMALL = []          # (page, pt, figure, text) - below the builder's floor, above the fail floor
CROP_TEXT = {}           # fingerprint -> [(native width pt, text size pt, text, file)] - deckspec.CROPTEXT
BACKUP_SMALL = []        # (page, text) - small text on a backup slide. Reported for reference, not counted as a failure
HLFIG_SMALL = []         # (page, text) - small text inside a pasted paper figure that points at where to look. Reported for reference
PASTED_TAG = " (inside a pasted figure)"


def crop_text(paths):
    """Reads the `croptext.tsv` file(s) left behind when cropping."""
    out = {}
    for p in paths:
        if not p or not os.path.isfile(p):
            continue
        with io.open(p, encoding="utf-8") as f_:
            for line in f_:
                parts = line.rstrip("\n").split("\t")
                if line.startswith("#") or len(parts) < 5:
                    continue
                try:
                    out.setdefault(parts[0], []).append((float(parts[1]), float(parts[2]),
                                                         parts[3], parts[4]))
                except ValueError:
                    continue
    return out


def check_page(page, n, min_size_pct, min_overlap, margin,
               min_contrast=0.0, figtext=None, labels=()):
    """The list of problems on one page. `(kind, description)`. `labels` = text placed
    on purpose on top of a figure."""
    R = page.rect
    H = R.height
    bad = []
    sp = spans(page)
    mf = math_flags(page)
    if len(mf) != len(sp):                     # if the order is off, don't exclude anything (play it safe)
        mf = [False] * len(sp)
    sf = span_fonts(page)
    if len(sf) != len(sp):
        sf = [None] * len(sp)
    _med = DECK_BODY[0]           # the deck-wide median size of deck-font text
    im = images(page)
    pix = page.get_pixmap(dpi=110) if min_contrast else None

    # (1) off the page / margin intrusion ─────────────────────────────
    #   These two are kept separate. Off the page is geometry and needs no
    #     tolerance (only 1px of render rounding), while the margin has its own
    #     standard: EBU R 95 sets the top/bottom/left/right 5% of a 16:9 screen as
    #     the graphics-safe zone. A talk gets recorded and projected, and goes through
    #     a projector's overscan and keystone correction, so this applies as-is.
    #   These two used to be lumped into one tolerance value. That made "ran off the
    #     page" hang on an arbitrary number, and "not enough margin" was not checked
    #     at all.
    #   A typesetter that silently drops an overflow is not caught here (beamer
    #     simply does not draw an overflowing line); that is what `outcheck` catches.
    #     Both are needed.
    safe = fitz.Rect(R.x0 + margin * R.width, R.y0 + margin * H,
                     R.x1 - margin * R.width, R.y1 - margin * H)
    for r, t, size, _ in sp:
        out = max(R.y0 - r.y0, r.y1 - R.y1, R.x0 - r.x0, r.x1 - R.x1)
        if out > EPS:
            bad.append(("off the page", "%.1fpt past the edge \"%s\"" % (out, t[:36])))
        elif not (r in safe):
            d = max(safe.y0 - r.y0, r.y1 - safe.y1,
                    safe.x0 - r.x0, r.x1 - safe.x1)
            bad.append(("margin", "%.1fpt outside the safe zone \"%s\"" % (d, t[:32])))
    for r in im:
        out = max(R.y0 - r.y0, r.y1 - R.y1, R.x0 - r.x0, r.x1 - R.x1)
        if out > EPS:
            bad.append(("off the page", "figure runs %.1fpt past the edge" % out))

    # (2) overlap ──────────────────────────────────────────────────────
    #   An attempt to skip within the same block missed real overlap. Two boxes
    #     occupying the same spot on one line is exactly what has to be caught, and
    #     that usually falls in the same block. Neighboring words on one line only
    #     touch and do not overlap, so the area threshold alone is enough - do not
    #     filter by block.
    for i in range(len(sp)):
        ri, ti, _, bi = sp[i]
        if mf[i]:
            continue
        for j in range(i + 1, len(sp)):
            rj, tj, _, bj = sp[j]
            if mf[j]:
                continue
            # If neither is a deck font, they are both text inside a pasted vector
            #   figure - two lines of a paper figure's label would otherwise be
            #   flagged as "overlap". What happens inside a figure is not the deck's to fix.
            if DECK_FONTS and sf[i] and sf[j] and sf[i] not in DECK_FONTS \
                    and sf[j] not in DECK_FONTS and sp[i][2] < _med and sp[j][2] < _med:
                continue
            # When an accent is pulled out attached to a longer fragment (e.g.
            #   "= max(ˆ"), it looks like it overlaps the letter under it ("v") - v̂ would
            #   get flagged as overlap. Excluded when one side contains a
            #   math mark and the other is three characters or fewer.
            if (any(c in MATH_MARKS for c in ti) and len(tj.strip()) <= 3) or \
                    (any(c in MATH_MARKS for c in tj) and len(ti.strip()) <= 3):
                continue
            f = overlap_frac(ri, rj)
            if f > min_overlap:
                bad.append(("overlap", "\"%s\" and \"%s\" overlap %.0f%%"
                            % (ti[:22], tj[:22], f * 100)))
    for r, t, _, _ in sp:
        _t = t.strip()
        if _t and any(_t in lab for lab in labels):
            continue                  # a figure.highlight label - on top of the figure is its rightful place
        for g in im:
            f = overlap_frac(r, g)
            if f > min_overlap:
                bad.append(("overlap", "\"%s\" sits %.0f%% on top of a figure" % (t[:26], f * 100)))

    # (3) text too small ────────────────────────────────────────────
    for k_, (r, t, size, _) in enumerate(sp):
        if mf[k_]:
            continue          # a script/accent is supposed to be small (a math accessory)
        if MARKER.fullmatch(t):
            continue          # a footnote marker is supposed to be small - explained above
        # A size equal to the floor passes. The floor is scriptsize (0.7x body),
        #   which is exactly the size the builder's `fine` and the page number use.
        #   Floating-point's last digit would otherwise flag them as "too small",
        #   so the recommended feature would always fail.
        if 100.0 * size / H < min_size_pct * 0.99:
            # a font the deck does not use = text inside a pasted vector figure.
            # marking it lets a figure with a called-out spot be excluded as reference
            _pasted = (bool(DECK_FONTS) and sf[k_] is not None and sf[k_] not in DECK_FONTS
                       and font_family(sf[k_]) not in {font_family(f_) for f_ in DECK_FONTS})
            bad.append(("small text", "%.1fpt = %.2f%% of page height (floor %.2f%%) \"%s\"%s"
                        % (size, 100.0 * size / H, min_size_pct, t[:26],
                           PASTED_TAG if _pasted else "")))

    # (3-b) text inside a figure ─────────────────────────────────
    #   Not in the PDF's text layer. The size the drawing side wrote down is measured
    #     with the same ruler. Without this, a 4.4pt diagram would pass every check.
    if figtext:
        # The builder warns below body size x 0.7 (7pt), and here anything below the
        #   page's median text size x 0.7 is counted as a failure. Otherwise the two
        #   drift apart silently into "build says small, fitcheck says fine". The
        #   failure floor is left as is (even a well-made deck often has
        #   figure text under 7pt - being strict here would let false positives erase
        #   the feature), and text that falls in between is counted separately, for
        #   reference.
        # The failure floor is never higher than the lowest value the builder
        #   draws (`deckspec.FIG_FLOOR_PT`). If the floor tracked the deck's median
        #   text size, the same figure text size could be reference in one deck
        #   but a failure in another - a figure the skill drew would fail the
        #   skill's own check depending on the deck. The builder already warns below
        #   7pt. If LaTeX shrinks the figure further, below the floor, that is a
        #   failure.
        _fig_lim = min(min_size_pct, 100.0 * deckspec.FIG_FLOOR_PT / H)
        _soft = 100.0 * deckspec.BODY_PT * 0.7 / H
        for key in image_prints(page.parent, page):
            for pt, s, nm in figtext.get(key, []):
                if _fig_lim * 0.99 <= 100.0 * pt / H < _soft * 0.99:
                    SOFT_SMALL.append((n, pt, nm, s))
                if 100.0 * pt / H < _fig_lim * 0.99:
                    bad.append(("small text",
                                "%.1fpt = %.2f%% of page height (floor %.2f%%) "
                                "inside figure %s \"%s\""
                                % (pt, 100.0 * pt / H, _fig_lim, nm,
                                   s[:22])))
    # Text inside a cropped paper figure - once cropped it is a PNG, so it is not
    #   in the text layer. Measured from the original size left behind at crop time,
    #   times the scale it was placed on the page at, to get its on-screen size
    #   (otherwise cropping would be a way to dodge this check).
    if CROP_TEXT:
        _fig_lim = min_size_pct
        _soft = 100.0 * deckspec.BODY_PT * 0.7 / H
        for info in page.get_images(full=True):
            try:
                bb = page.get_image_bbox(info)
                # Naming this `pix` would overwrite the page render the contrast
                #   check below uses, which would get a title flagged as
                #   "same color as the background" by mistake
                _px = fitz.Pixmap(page.parent, info[0])
                if _px.alpha or (_px.colorspace and _px.n > 3):
                    _px = fitz.Pixmap(fitz.csRGB, _px)
                im = Image.frombytes("RGB" if _px.n >= 3 else "L", (_px.width, _px.height),
                                     _px.samples)
                key = fingerprint(im)
            except Exception:                              # noqa: BLE001
                continue
            _small = []
            for native_w, size, s, nm in CROP_TEXT.get(key, []):
                pt = size * bb.width / max(1e-6, native_w)
                if _fig_lim * 0.99 <= 100.0 * pt / H < _soft * 0.99:
                    SOFT_SMALL.append((n, pt, nm, s))
                if 100.0 * pt / H < _fig_lim * 0.99:
                    _small.append((pt, s, nm))
            if _small:
                pt, s, nm = min(_small)
                bad.append(("small text",
                            "%.1fpt = %.2f%% of page height (floor %.2f%%) inside cropped "
                            "figure %s \"%s\" and %d more - "
                            "the paper figure's text stays small. Enlarge one panel, or redraw it"
                            % (pt, 100.0 * pt / H, _fig_lim, nm, s[:22], len(_small) - 1)))

    # (4) contrast (WCAG) ──────────────────────────────────────────────
    #   Measuring the color agrees with human judgment better than asking a
    #     model. It is also one of the few items with an actual standard - no need
    #     to invent one.
    if pix is not None:
        for r, t, size, col in sp:
            if MARKER.fullmatch(t):
                continue
            ink = ((col >> 16) & 255, (col >> 8) & 255, col & 255)
            paper = paper_under(pix, r, R, ink)
            if paper is None:
                continue
            # A standout slide's page number is deliberately painted the
            #   background color by the theme - it is not meant to be read. Only a
            #   page number where the background and text are the same color is
            #   excluded.
            if paper == ink and FRAME_NO.fullmatch(t.strip()):
                continue
            c = contrast(ink, paper)
            if c < min_contrast:
                # to two decimal places - 4.47 was printing as `4.5 (floor 4.5)`.
                bad.append(("low contrast", "%.2f:1 (floor %.1f) bg %s / text %s \"%s\""
                            % (c, min_contrast, _hex(paper), _hex(ink), t[:24])))
    return bad


def empty_para_pt(p, pt):
    """The font size that sets an empty paragraph's line height.

    PowerPoint bases an empty paragraph's height not on a run but on the size of
      the end-of-paragraph mark (`endParaRPr`), falling back to a default 18pt if
      there is none. Otherwise the row of a 15pt table with an empty cell would grow
      to 0.40in, and the note below the table would overlap the last row, since
      measuring by run size misses it (confirmed against the row height PowerPoint
      itself reported).
    """
    if "".join(x.text for x in p.runs):
        return pt
    from pptx.oxml.ns import qn
    e = p._p.find(qn("a:endParaRPr"))
    if e is not None and e.get("sz"):
        return int(e.get("sz")) / 100.0
    return 18.0


def ink_rect(sh, r):
    """The rectangle the text actually fills inside a text box (inches). The text,
    not the frame.

    For each paragraph, height comes from font size x wrapped line count, placed by
    vertical alignment. Width is the longest line's measured text width (by font),
    placed by horizontal alignment.
    """
    from pptx.util import Emu
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    tf = sh.text_frame

    def _in(v, d):
        return Emu(v).inches if v is not None else d
    li, ri = _in(tf.margin_left, 0.1), _in(tf.margin_right, 0.1)
    ti, bi = _in(tf.margin_top, 0.05), _in(tf.margin_bottom, 0.05)
    w_in = max(0.2, (r[2] - r[0]) - li - ri)
    h = 0.0
    wmax = 0.0
    align = None
    first, last = None, None
    for p in tf.paragraphs:
        s = "".join(x.text for x in p.runs) or p.text or ""
        pt = None
        for x in p.runs:
            if x.font.size is not None:
                pt = x.font.size.pt
                break
        if pt is None and p.font.size is not None:
            pt = p.font.size.pt
        pt = empty_para_pt(p, pt or 18.0)
        em = pt / 72.0 * 0.50
        # Measured by font and wrapped word by word - an average-width estimate
        #   would carry the same error as the builder and miss overlapping text.
        _bold = any(x.font.bold for x in p.runs)
        _fn = next((x.font.name for x in p.runs if x.font.name), None) or "Arial"
        n = (deckspec.wrap_count(s, w_in, pt, _bold, _fn)
             if tf.word_wrap is not False else 1)
        # Line spacing is counted. The builder uses 1.35-1.45 for bullets, but
        #   measuring at 1.2 would let an overflowing bullet covering the block below
        #   still read "overlap 0".
        ls = p.line_spacing if isinstance(p.line_spacing, float) else None
        # PowerPoint's one line-slot = font size x 1.2 x line spacing, and the
        #   first line too (measured from a render: 20pt, 1.35 -> one line is
        #   0.458in). Counting the first line at 1.2 alone would leave a seven-line
        #   bullet's end touching the next block undetected. Instead, the range the
        #   text actually paints is used - the spacing above the
        #   first line and the empty share below the last line are not painted
        #   (guards against false positives from large numbers).
        L = pt * 1.2 * (ls or 1.0) / 72.0
        if first is None:
            first = (L, pt)
        last = pt
        h += n * L
        wmax = max(wmax, min(w_in, deckspec.text_width_in(s, pt, _bold, _fn)))
        align = align or p.alignment
    # Only a box that shrinks text to fit the frame is confined to the frame.
    #   Otherwise PowerPoint does not cut off overflowing text - text overflowing
    #   and covering the box below is exactly what this check looks for.
    if tf.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE:
        h = min(h, (r[3] - r[1]) - ti - bi)
    anc = tf.vertical_anchor
    top = r[1] + ti
    if anc == MSO_ANCHOR.MIDDLE:
        top = (r[1] + r[3]) / 2.0 - h / 2.0
    elif anc == MSO_ANCHOR.BOTTOM:
        top = r[3] - bi - h
    left = r[0] + li
    if align == PP_ALIGN.CENTER:
        left = (r[0] + r[2]) / 2.0 - wmax / 2.0
    elif align == PP_ALIGN.RIGHT:
        left = r[2] - ri - wmax
    # the painted range: from the first line's cap top (0.97em above the line-slot's
    # bottom) to 0.15em below the last line's baseline
    first = first or (0.0, 0.0)
    y0 = top + max(0.0, first[0] - 0.97 * first[1] / 72.0)
    y1 = max(y0 + 0.01, top + h - 0.15 * (last or 0.0) / 72.0)
    return (left, y0, left + wmax, y1)


def split_tokens(sh, r):
    """A single chunk (no spaces) that is wider than the box - PowerPoint breaks the
    line in the middle of it.

    An impact slide's "93,817.45M" could print as "93,81 / 7.45M", but a check that
      measures by line count says nothing since the two lines fit in the box.
      A split number reads as a different number.
    """
    from pptx.util import Emu
    tf = sh.text_frame
    if tf.word_wrap is False:
        return []
    li = Emu(tf.margin_left).inches if tf.margin_left is not None else 0.1
    ri = Emu(tf.margin_right).inches if tf.margin_right is not None else 0.1
    w_in = (r[2] - r[0]) - li - ri
    out = []
    for p in tf.paragraphs:
        s = "".join(x.text for x in p.runs)
        pt = next((x.font.size.pt for x in p.runs if x.font.size), None) or 18.0
        bold = any(x.font.bold for x in p.runs)
        fn = next((x.font.name for x in p.runs if x.font.name), None) or "Arial"
        for tok in s.split():
            if len(tok) > 1 and deckspec.text_width_in(tok, pt, bold, fn) > w_in + 0.01:
                out.append(tok)
        # A short value containing a number ("41.7 ms") is still split even when
        #   it breaks into two lines - each word fits on its own, so the check above
        #   would otherwise miss it.
        t_ = s.strip()
        if (" " in t_ and len(t_) <= deckspec.BIG_VALUE_CHARS and re.search(r"\d", t_)
                and t_ not in out and deckspec.wrap_count(t_, w_in, pt, bold, fn) > 1):
            out.append(t_)
    return out


def table_rendered_h(sh):
    """The height (inches) a table fills on screen - measured by font per cell, counting
    wrapped lines."""
    from pptx.util import Emu
    tbl = sh.table
    widths = [Emu(c.width).inches for c in tbl.columns]
    tot = 0.0
    for row in tbl.rows:
        need = Emu(row.height).inches
        for ci, cell in enumerate(row.cells):
            if ci >= len(widths) or cell.is_spanned:
                continue
            w = widths[ci] - (Emu(cell.margin_left).inches + Emu(cell.margin_right).inches)
            if cell.is_merge_origin:
                w = sum(widths[ci:ci + cell.span_width]) - 0.2
            h = 0.0
            for p in cell.text_frame.paragraphs:
                s = "".join(x.text for x in p.runs) or ""
                pt = empty_para_pt(
                    p, next((x.font.size.pt for x in p.runs if x.font.size), 12.0))
                bold = any(x.font.bold for x in p.runs)
                fn = next((x.font.name for x in p.runs if x.font.name), None) or "Arial"
                h += deckspec.wrap_count(s, max(0.2, w), pt, bold, fn) * pt * 1.2 / 72.0
            h += Emu(cell.margin_top).inches + Emu(cell.margin_bottom).inches
            need = max(need, h)
        tot += need
    return tot


def pptx_boxes(path):
    """A PPTX's shapes as (kind, rect, text). Inches."""
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(path)
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    out = []
    for n, sl in enumerate(prs.slides, 1):
        got = []
        for sh in sl.shapes:
            if sh.top is None or sh.left is None:
                continue
            r = (Emu(sh.left).inches, Emu(sh.top).inches,
                 Emu(sh.left).inches + Emu(sh.width).inches,
                 Emu(sh.top).inches + Emu(sh.height).inches)
            if "PICTURE" in str(sh.shape_type):
                got.append(("pic", r, ""))
            elif getattr(sh, "has_table", False) and sh.has_table:
                # A table also takes up space - without this, text placed below a
                #   table printing on top of the table would go undetected.
                # Height is re-measured from wrapped lines, not the file's value.
                #   PowerPoint only grows a wrapped row on screen, so trusting the
                #   builder's estimate as-is would miss the same error.
                r = (r[0], r[1], r[2], r[1] + table_rendered_h(sh))
                got.append(("tbl", r, "table"))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                # a highlight label's rightful place is on top of the figure - excluded only from figure-overlap
                kind = "hl" if sh.name == HIGHLIGHT_TAG else "txt"
                got.append((kind, r, sh.text_frame.text.strip()))
                # a page number sits in its own corner - its area overlap is checked, but it is excluded from the
                #   "touching" check. Otherwise a bullet block's bounding box reaching into the corner
                #   reads as a false positive.
                got.append(("inkpg" if sh.name == PAGENUM_TAG else "ink",
                            ink_rect(sh, r), sh.text_frame.text.strip()))
                for tok in split_tokens(sh, r):
                    got.append(("split", r, tok))
        out.append((n, got))
    return out, W, H


def check_pptx(path, min_overlap_in2=0.05, eps=0.02):
    """Spots in a PPTX where text covers a figure or runs off the screen.

    `deckcheck` section F counts the set of text, section G counts the slide count.
      Neither one looks at where things are placed. Text can run across a
      figure on several slides with no check raising a flag, which only shows up by
      rendering the PPTX and looking at it. The machine in the room is PowerPoint.
      Measuring only the PDF deck is not enough.
    """
    bad = []
    pages, W, H = pptx_boxes(path)
    for n, shapes in pages:
        pics = [r for k, r, _ in shapes if k in ("pic", "tbl")]
        _out = set()
        for k, r, s in shapes:
            if k == "split":
                continue                  # checked separately below - its position is the same as that text box
            # Measuring only the frame misses overflowing text that ran off the
            #   screen - PowerPoint does not clip overflowing text. The space the text
            #   fills (`ink`) is measured too.
            out = max(-r[0], -r[1], r[2] - W, r[3] - H)
            if out > eps and (k, s) not in _out and ("txt", s) not in _out:
                _out.add((k, s))
                bad.append((n, "off the page",
                            "%s runs %.2f in past the edge \"%s\""
                            % ("figure" if k == "pic" else "text", out, s[:26])))
            if k != "txt":
                continue
            for g in pics:
                w = min(r[2], g[2]) - max(r[0], g[0])
                h = min(r[3], g[3]) - max(r[1], g[1])
                a = max(0.0, w) * max(0.0, h)
                if a > min_overlap_in2:
                    bad.append((n, "overlap",
                                "\"%s\" covers %.2f in² of a figure/table" % (s[:26], a)))
        for k, r, s in shapes:
            if k == "split":
                bad.append((n, "split", "\"%s\" is wider than the box, so the line breaks in the "
                                         "middle of it - shrink the text or widen the box" % s[:26]))
        # A label placed on top of a figure (highlight) has its rightful place on
        #   top of the figure, so it is excluded from figure overlap - but that also
        #   excludes it from table overlap unless checked separately here. Otherwise a
        #   label overflowing the figure onto an adjacent table's header would still
        #   read 0.
        tbls = [r for k, r, _ in shapes if k == "tbl"]
        for k, r, s in shapes:
            if k != "hl":
                continue
            for g in tbls:
                w = min(r[2], g[2]) - max(r[0], g[0])
                h = min(r[3], g[3]) - max(r[1], g[1])
                if max(0.0, w) * max(0.0, h) > min_overlap_in2:
                    bad.append((n, "overlap", "label \"%s\" on the figure covers %.2f in² of the "
                                         "adjacent table - shrink the label or move it with "
                                         "`label_at`" % (s[:24], w * h)))
        # Text against text - only the space the text fills is checked against
        #   itself. Frames always overlap each other.
        inks = [(r, s) for k, r, s in shapes if k in ("ink", "inkpg")]
        body_inks = [(r, s) for k, r, s in shapes if k == "ink"]
        for i, (r1, s1) in enumerate(inks):
            for r2, s2 in inks[i + 1:]:
                w = min(r1[2], r2[2]) - max(r1[0], r2[0])
                h = min(r1[3], r2[3]) - max(r1[1], r2[1])
                if w <= 0 or h <= 0:
                    continue
                # Since these are painted ranges, any overlap at all means the
                #   text touches - a threshold like "25% of the smaller box" would let a
                #   whole line touching through undetected.
                if w * h > min_overlap_in2 and h > 0.02:
                    bad.append((n, "overlap",
                                "\"%s\" and \"%s\" overlap by %.2f in²"
                                % (s1[:18], s2[:18], w * h)))
        # Touching is also checked - two blocks of text stacked with no gap
        #   between them read as one block on screen. Overlap area alone reads 0, so
        #   it would otherwise miss a bullet's last line touching the next block.
        for i, (r1, s1) in enumerate(body_inks):
            for r2, s2 in body_inks[i + 1:]:
                w = min(r1[2], r2[2]) - max(r1[0], r2[0])
                narrow = min(r1[2] - r1[0], r2[2] - r2[0])
                if w < 0.5 * narrow or w <= 0:
                    continue
                a_, b_ = (r1, r2) if r1[1] <= r2[1] else (r2, r1)
                gap = b_[1] - a_[3]
                if -0.02 <= gap < 0.05:
                    bad.append((n, "overlap", "\"%s\" has \"%s\" directly below it with no gap (%.2f in)"
                                % ((s1 if a_ is r1 else s2)[:18], (s2 if a_ is r1 else s1)[:18],
                                   gap)))
    return bad


# The function that reads overflow is one single one - `deckspec.log_overfull`.
#   If the feedback side (build.py) and the reporting side (here) used different
#   rules, one would still call it overflowing after the other one fixed it.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckspec import log_overfull, HIGHLIGHT_TAG, HIGHLIGHT_LIST  # noqa: E402
import deckspec  # noqa: E402
try:
    from build_pptx import PAGENUM_TAG  # noqa: E402  (the page-number box's name - the builder is the single source)
except Exception:                       # the PDF check still runs even without python-pptx
    PAGENUM_TAG = "p2t-pagenum"


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Does the rendered deck fit in its place - overflow, overlap, small text")
    ap.add_argument("pdf")
    ap.add_argument("--viewing-ratio", type=float, default=None, metavar="N",
                    help="how many times the screen height away is the back row "
                         "(e.g. 6 = a classroom). If given, uses AVIXA BDM's absolute "
                         "floor. If not given, only flags text noticeably smaller "
                         "than the body")
    ap.add_argument("--min-height", type=float, default=None, metavar="PCT",
                    help="give the text-size floor directly, as %% of page height. "
                         "If given, overrides --viewing-ratio")
    ap.add_argument("--cap-ratio", type=float, default=CAP_RATIO, metavar="F",
                    help="cap height / font size (default %.2f). Varies by font, "
                         "0.62-0.72" % CAP_RATIO)
    ap.add_argument("--overlap", type=float, default=0.18, metavar="FRAC",
                    help="flag overlap greater than this (fraction of the smaller "
                         "area, default 0.18)")
    ap.add_argument("--margin", type=float, default=0.0, metavar="FRAC",
                    help="this much inward from the edge is the safe zone. Use 0.05 "
                         "(EBU R 95) for a recorded/broadcast talk. Default is to not "
                         "check it - a theme's title band hits every page")
    ap.add_argument("--contrast", type=float, default=4.5, metavar="RATIO",
                    help="floor for the contrast ratio between text and background "
                         "(default 4.5 = WCAG 2.2 AA). 0 disables the check")
    ap.add_argument("--figtext", default=None, metavar="TSV",
                    help="figs/textsize.tsv left by `build_figs`. Text inside a "
                         "figure is not in the PDF's text layer, so without this "
                         "not a single piece of it is counted")
    ap.add_argument("--pptx", default=None, metavar="FILE",
                    help="measure the PPTX too - does text cover a figure or run off "
                         "the screen. The machine in the room is PowerPoint")
    ap.add_argument("--show", type=int, default=6, help="how many to show per kind")
    a = ap.parse_args(argv)

    d = fitz.open(a.pdf)
    # If not given, looks for one nearby - if it only worked "when given", one would just skip giving it.
    ft_path = a.figtext
    if ft_path is None:
        for cand in ("figs/textsize.tsv", "textsize.tsv"):
            p = os.path.join(os.path.dirname(os.path.abspath(a.pdf)), cand)
            if os.path.exists(p):
                ft_path = p
                break
    figtext = fig_text(ft_path)
    # a cropped figure's file appears next to the original (usually the working
    # folder's figs/, while the deck is in out/) - looked for up to two levels up
    import glob as _glob
    _pd = os.path.dirname(os.path.abspath(a.pdf))
    _ct = set()
    for _base in (_pd, os.path.dirname(_pd)):
        for _pat in (deckspec.CROPTEXT, os.path.join("*", deckspec.CROPTEXT),
                     os.path.join("*", "*", deckspec.CROPTEXT)):
            _ct |= set(_glob.glob(os.path.join(_base, _pat)))
    CROP_TEXT.clear()
    CROP_TEXT.update(crop_text(sorted(_ct)))
    _bkp = os.path.join(os.path.dirname(os.path.abspath(a.pdf)), "figs", deckspec.BACKUP_LIST)
    backup_pages = set()
    if os.path.exists(_bkp):
        with io.open(_bkp, encoding="utf-8") as f_:
            backup_pages = {int(x) for x in f_.read().split() if x.strip().isdigit()}
    BACKUP_SMALL.clear()
    HLFIG_SMALL.clear()
    _hfp = os.path.join(os.path.dirname(os.path.abspath(a.pdf)), "figs", deckspec.HLFIG_LIST)
    hlfig_pages = set()
    if os.path.exists(_hfp):
        with io.open(_hfp, encoding="utf-8") as f_:
            hlfig_pages = {int(x) for x in f_.read().split() if x.strip().isdigit()}
    _hlp = os.path.join(os.path.dirname(os.path.abspath(a.pdf)), "figs", HIGHLIGHT_LIST)
    labels = []
    if os.path.exists(_hlp):
        with io.open(_hlp, encoding="utf-8") as f_:
            labels = [x.strip() for x in f_ if x.strip()]
    print("=" * 78)
    print("layout check - %s (%d page(s), %.0f x %.0f pt)"
          % (a.pdf, len(d), d[0].rect.width, d[0].rect.height))
    print("=" * 78)
    if figtext:
        print("   also counting %d piece(s) of text inside figures (%s)"
              % (sum(len(v) for v in figtext.values()),
                 os.path.basename(ft_path)))
    else:
        print("   Text inside figures is not counted - it is not in the PDF's text layer.")
        print("     Run `build_figs` to make figs/textsize.tsv and pass it with "
              "`--figtext`.")

    kinds = {}
    tiny_pages = set()
    caps = []
    # If no distance is given, the body is the reference. Only text under 0.7x the
    #   body's median (scriptsize) counts as small - defaulting to an absolute
    #   standard (6x) would flag 10pt body text on every page, leaving no signal.
    _all = sorted(100.0 * s_[2] / d[0].rect.height for p_ in d for s_ in spans(p_))
    _body = _all[len(_all) // 2] if _all else 3.5
    DECK_FONTS.clear()
    DECK_FONTS.update(deck_fonts(d, hlfig_pages))
    _ds = sorted(s_[2] for p_ in d for s_, f_ in zip(spans(p_), span_fonts(p_))
                 if f_ in DECK_FONTS)
    DECK_BODY[0] = _ds[len(_ds) // 2] if _ds else 0.0
    for i, p in enumerate(d, 1):
        cr, measured = cap_ratio_of(p)
        cr = a.cap_ratio if a.cap_ratio != CAP_RATIO else cr
        if a.min_height is not None:
            lim = a.min_height
        elif a.viewing_ratio is not None:
            lim = min_cap_pct(a.viewing_ratio) / cr
        else:
            lim = _body * 0.7   # LaTeX's size ladder: footnotesize 0.8, scriptsize 0.7
            # The floor must not drop along with the deck. Otherwise a table's text
            #   shrinking quietly would shrink the deck-wide median with it, dropping
            #   the floor to 1.21% and reporting "small text on 0 pages". 0.6x the body
            #   size (6pt for a 10pt body) is small in any deck.
            lim = max(lim, 100.0 * deckspec.BODY_PT * 0.6 / p.rect.height)
        caps.append((cr, measured, lim))
        for kind, msg in check_page(p, i, lim, a.overlap, a.margin,
                                    a.contrast, figtext, labels):
            # Small text on a backup slide is reference only. Otherwise a slide opened
            #   only if a question comes up would have no escape from a paper figure
            #   that cannot be redrawn, and would always exit 1. Overflow/overlap
            #   are still counted as-is.
            if kind == "small text" and i in backup_pages:
                BACKUP_SMALL.append((i, msg))
                continue
            # Tick-mark text in a pasted paper figure with a called-out spot is
            #   also reference only. Otherwise, since planning's §plots says "if a
            #   result only exists in a figure, paste it and point at it with
            #   highlight", that exact placement would be counted as a failure - and
            #   only for vector figures, so swapping the same figure to a PNG would
            #   pass, nudging people toward raster for the wrong reason. Deck text is
            #   still counted as-is.
            if kind == "small text" and i in hlfig_pages and (
                    msg.endswith(PASTED_TAG) or "cropped figure" in msg):
                HLFIG_SMALL.append((i, msg))
                continue
            kinds.setdefault(kind, []).append((i, msg))
            if kind == "small text":
                tiny_pages.add(i)

    # Where the floor came from is always printed alongside it. If it looks like
    #   an arbitrary number, a person ignores it, and an ignored check is the same as
    #   no check.
    cr = caps[0][0] if caps else CAP_RATIO
    measured = any(m for _, m, _ in caps)
    lim = caps[0][2] if caps else 0.0
    sizes = sorted(100.0 * s[2] / d[0].rect.height for p in d for s in spans(p))
    if sizes:
        print("   text size (as %% of page height)  min %.2f%% - median %.2f%% - max %.2f%%"
              % (sizes[0], sizes[len(sizes) // 2], sizes[-1]))
        if a.min_height is None and a.viewing_ratio is None:
            if lim > _body * 0.7 + 1e-6:
                print("   floor %.2f%% = 0.6x the %.0fpt body (%.0fpt) - higher than 0.7x "
                      "the median (%.2f%%), so this one is used (the floor does not drop "
                      "along with the deck)"
                      % (lim, deckspec.BODY_PT, deckspec.BODY_PT * 0.6, _body))
            else:
                print("   floor %.2f%% = 0.7x the body's median (%.2f%%) (scriptsize) - if "
                      "you know the venue's viewing distance, give "
                      "`--viewing-ratio`" % (lim, _body))
        elif a.min_height is None:
            print("   floor %.2f%% = back-row distance %.1fx (AVIXA BDM: cap height %.2f%%)"
                  " / cap ratio %.2f%s"
                  % (lim, a.viewing_ratio, min_cap_pct(a.viewing_ratio), cr,
                     " (measured)" if measured else " (estimated)"))
            # How far the smallest text is still legible - given as a fact, not a verdict.
            reach = sizes[0] * cr * AVIXA_ACUITY / 100.0
            print("   -> the smallest text is legible up to %.1fx the screen height." % reach)
        else:
            print("   floor %.2f%% (given directly)" % lim)
    print()

    # Pages where the body's own placeholder overflowed. Even while still inside
    #   the page, it eats into the page-number/footer area.
    #   LaTeX already knows this and wrote it to the log - nobody just read the log.
    _ov = log_overfull(a.pdf)
    if _ov is None:
        print("   (no .log next to it, so body overflow could not be measured - "
              "give a PDF compiled in the same place)")
    for n, pt in (_ov or []):
        kinds.setdefault("body overflow", []).append(
            (n, "the body is %.1fpt taller than its space - it eats into the page-number "
                "area. Trim it (start with `fine`/`foot`) or split the slide" % pt))

    bad = 0
    for kind in ("off the page", "body overflow", "margin", "overlap",
                 "small text", "low contrast"):
        hits = kinds.get(kind, [])
        # The unit counted is pages to go fix. It is normal for the same cause to
        #   hit 80 times on one page, so counting by fragment would give "1105
        #   problems", a number that says nothing. "14 pages" is a to-do list.
        pages = sorted({n for n, _ in hits})
        bad += len(pages)
        print("   %s %-13s %d page(s)%s"
              % ("! " if hits else "  ", kind, len(pages),
                 "  (%d fragment(s))" % len(hits) if len(hits) != len(pages) else ""))
        # If one kind has hundreds of hits, printing every line means nobody reads
        #   it. The same cause hitting a page dozens of times is normal, so only the
        #   worst one per page is kept. A check with a list so long it gets ignored is
        #   the same as no check.
        by_page = {}
        for n, msg in hits:
            by_page.setdefault(n, msg)
        rows = sorted(by_page.items())
        for n, msg in rows[:a.show]:
            extra = sum(1 for x, _ in hits if x == n) - 1
            print("        page %2d  %s%s"
                  % (n, msg, "   (%d more on this page)" % extra if extra else ""))
        if len(rows) > a.show:
            print("        ... and %d more page(s)" % (len(rows) - a.show))
    d.close()

    if a.pptx:
        print()
        print("   " + "-" * 46)
        try:
            pb = check_pptx(a.pptx)
        except ImportError:
            print("   could not measure the PPTX - `pip install python-pptx`")
            pb = []
        pages_ = sorted({n for n, _, _ in pb})
        print("   %s PPTX geometry  %d slide(s)%s"
              % ("! " if pb else "  ", len(pages_),
                 "  (%d hit(s))" % len(pb) if len(pb) != len(pages_) else ""))
        for n, kind, msg in pb[:a.show]:
            print("        slide %2d  %s  %s" % (n, kind, msg))
        if len(pb) > a.show:
            print("        ... and %d more" % (len(pb) - a.show))
        if pb:
            print("     The PDF deck constrains both width and height; the PPTX side used to "
                  "constrain only width.")
            print("       `deckcheck` F/G do not look at where things are placed -")
            print("       they only compare the set of text and the slide count. Both can "
                  "pass while things overlap.")
        bad += len(pages_)

    print()
    print("   pages to fix: %d" % bad)
    if BACKUP_SMALL:
        _bp = sorted({p_ for p_, _m in BACKUP_SMALL})
        print("   (info) small text on backup slides: %d - not counted as a failure "
              "(not opened during the talk): page(s) %s. "
              "Leave it so it can be zoomed in on if a question opens it"
              % (len(BACKUP_SMALL), ",".join(str(x) for x in _bp)))
    if HLFIG_SMALL:
        _hp = sorted({p_ for p_, _m in HLFIG_SMALL})
        print("   (info) small text inside a paper figure that is either called out "
              "or shows several panels at once: %d - not counted as a failure: page(s) %s. "
              "The audience cannot read tick marks. Put values worth reading in a highlight "
              "label or spell them out in a side pane"
              % (len(HLFIG_SMALL), ",".join(str(x) for x in _hp)))
    if SOFT_SMALL:
        _pg = sorted({x[0] for x in SOFT_SMALL})
        print("   (info) %d piece(s) of text inside figures are below the builder's "
              "%.0fpt floor - not a failure (still above the page's own text floor): "
              "page(s) %s. Same fact as build's own warning."
              % (len(SOFT_SMALL), deckspec.BODY_PT * 0.7, ",".join(str(p) for p in _pg)))
    if kinds.get("small text"):
        print("   small text - text noticeably smaller than the body. This is what the "
              "audience misses first.")
        print("     For deck text, there are two ways out: cut content, or split the slide.")
        _pp = sorted({p_ for p_, m_ in kinds["small text"]
                      if m_.endswith(PASTED_TAG) or "cropped figure" in m_})
        if _pp:
            # Text in a pasted paper figure cannot be trimmed, so a separate path is
            #   given for it below.
            print("     Text inside a pasted paper figure (page(s) %s): enlarge one panel "
                  "with `crop`, redraw it, or point at the spot worth seeing with "
                  "`highlight` - a called-out figure's tick marks are only counted as "
                  "reference."
                  % ",".join(str(x) for x in _pp))
        print("     If you know the venue's back-row distance, set an absolute floor with "
              "`--viewing-ratio`.")
    if kinds.get("margin"):
        print("   margin - checked for a recorded/broadcast talk (EBU R 95, 5% top/bottom/left/right).")
        print("     The theme decides this value, so it should be the same on every slide.")
    if not bad:
        print("   Nothing overflows, overlaps, is too small, or is too low-contrast.")
    print("   (This only tells you where to point your eyes. It cannot see an ugly layout.)")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8")
    sys.exit(main())
